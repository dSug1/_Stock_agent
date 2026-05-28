"""Tests for the docx → csv converter (Module 0 preprocessing)."""
from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest
from bs4 import BeautifulSoup

HERE = Path(__file__).resolve()
PROJECT_ROOT = HERE.parents[1]
SRC = PROJECT_ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from docx_converter.convert import (  # noqa: E402
    COLUMN_MAP,
    EXPECTED_DOCX_COLS,
    _cell_blurred_or_text,
    _cell_text,
    _parse_price_history,
    _parse_sentiment,
    auto_convert_directory,
    convert,
    row_to_dict,
)


# =====================================================================
# Unit tests on cell-level extractors
# =====================================================================

def _cell(html: str):
    return BeautifulSoup(f"<table><tr><td>{html}</td></tr></table>", "html.parser").find("td")


def test_blurred_or_text_prefers_attribute():
    c = _cell('<div blurred-text="canonical">visible noise</div>')
    assert _cell_blurred_or_text(c) == "canonical"


def test_blurred_or_text_falls_back_to_text_when_no_attribute():
    c = _cell("just plain text")
    assert _cell_blurred_or_text(c) == "just plain text"


def test_text_extractor_collapses_whitespace():
    c = _cell("<span>foo</span>  <span>bar</span>")
    assert _cell_text(c) == "foo bar"


def test_price_history_keeps_only_even_indices():
    c = _cell('<div blurred-text="100.5,1700000000,101.0,1700086400,102.5,1700172800"></div>')
    assert _parse_price_history(c) == "100.5; 101.0; 102.5"


def test_price_history_empty_when_no_blurred_text():
    c = _cell("plain text only")
    assert _parse_price_history(c) == ""


def test_sentiment_parses_three_percentages():
    c = _cell("Community 67% 22% 11% Some Drug Name How are you feeling about this catalyst?")
    assert _parse_sentiment(c) == "Bull 67% / Neutral 22% / Bear 11%"


def test_sentiment_handles_no_votes_with_dashes():
    c = _cell("Community -% -% -% Drug name etc.")
    assert _parse_sentiment(c) == "Bull -% / Neutral -% / Bear -%"


def test_sentiment_defaults_to_blanks_when_no_match():
    c = _cell("totally unrelated content")
    assert _parse_sentiment(c) == "Bull -% / Neutral -% / Bear -%"


# =====================================================================
# COLUMN_MAP invariants
# =====================================================================

def test_column_map_has_19_columns():
    assert len(COLUMN_MAP) == 19


def test_column_map_drops_docx_index_9_options():
    indices = {idx for _, idx, _ in COLUMN_MAP}
    assert 9 not in indices, "docx col 9 (Options) must be dropped"


def test_column_map_uses_only_valid_modes():
    valid_modes = {"blurred_or_text", "text", "price_history", "sentiment"}
    for name, _, mode in COLUMN_MAP:
        assert mode in valid_modes, f"{name} has invalid mode {mode!r}"


def test_column_map_drug_and_catalyst_use_blurred_or_text():
    """Drug visible text appends FDA badge labels; Catalyst visible text
    is truncated with '… read more'. Both MUST use blurred_or_text."""
    by_name = {name: mode for name, _, mode in COLUMN_MAP}
    assert by_name["Drug"] == "blurred_or_text"
    assert by_name["Catalyst"] == "blurred_or_text"


# =====================================================================
# row_to_dict — synthetic 20-cell row
# =====================================================================

def _synth_row_html() -> str:
    """A minimal but realistic 20-cell row matching the BPC HTML shape."""
    cells = [
        '<div blurred-text="ABC"><a href="/company/ABC">ABC</a></div>',                          # 0 Ticker
        '<div blurred-text="ABC Co.">ABC Co.</div>',                                              # 1 Name
        '<div blurred-text="42.1234"><div class="price">$42.12</div> <small>+1.00 +2.4%</small></div>',  # 2 Price
        '<div blurred-text="10.0,1000,11.0,2000,12.0,3000"></div>',                               # 3 30-day
        '<div blurred-text="MyDrug (XYZ-1)"><strong>MyDrug (XYZ-1)</strong> <span>FTD</span></div>',  # 4 Drug
        'NCT01234567',                                                                            # 5 NCT
        'oncology indication',                                                                    # 6 Indication
        '<div blurred-text="phase2">phase 2</div>',                                               # 7 Stage
        '<div blurred-text="Ongoing">Ongoing</div>',                                              # 8 Status
        '<a>View</a>',                                                                            # 9 Options (DROPPED)
        '<div blurred-text="Topline Data">Topline Data</div>',                                    # 10 Next Catalyst
        '<div blurred-text="2026-07-15">15/07/2026 ET</div>',                                     # 11 Catalyst Date
        '<div blurred-text="Topline data expected Q3 2026">Topline data expected Q3 2026</div>',  # 12 Catalyst
        '',                                                                                       # 13 Conference (blank)
        '<div blurred-text="65.5">65.5</div>',                                                    # 14 Historical LOA
        '<div blurred-text="80.2">80.2</div>',                                                    # 15 Historical POP
        'Community 50% 30% 20% MyDrug How are you feeling about this catalyst?',                 # 16 Sentiment
        '<div blurred-text="1234567890">1.23B</div>',                                             # 17 Market Cap
        '<div blurred-text="2026-05-27 09:00:00">27/05/2026 ET</div>',                            # 18 Last Updated
        '<div blurred-text="98765432">98.7M</div>',                                               # 19 No Of Shares
    ]
    return "<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>"


def test_row_to_dict_produces_19_keys():
    soup = BeautifulSoup(_synth_row_html(), "html.parser")
    tds = soup.find("tr").find_all("td")
    assert len(tds) == EXPECTED_DOCX_COLS
    row = row_to_dict(tds)
    assert set(row.keys()) == {name for name, _, _ in COLUMN_MAP}


def test_row_to_dict_drug_strips_fda_badges():
    soup = BeautifulSoup(_synth_row_html(), "html.parser")
    tds = soup.find("tr").find_all("td")
    row = row_to_dict(tds)
    assert row["Drug"] == "MyDrug (XYZ-1)"        # NOT "MyDrug (XYZ-1) FTD"


def test_row_to_dict_dates_are_iso():
    soup = BeautifulSoup(_synth_row_html(), "html.parser")
    tds = soup.find("tr").find_all("td")
    row = row_to_dict(tds)
    assert row["Catalyst Date"] == "2026-07-15"
    assert row["Last Updated"] == "2026-05-27 09:00:00"


def test_row_to_dict_market_cap_is_integer_string():
    soup = BeautifulSoup(_synth_row_html(), "html.parser")
    tds = soup.find("tr").find_all("td")
    row = row_to_dict(tds)
    assert row["Market Cap"] == "1234567890"      # NOT "1.23B"
    assert row["No Of Shares"] == "98765432"


def test_row_to_dict_price_history_is_semicolon_joined():
    soup = BeautifulSoup(_synth_row_html(), "html.parser")
    tds = soup.find("tr").find_all("td")
    row = row_to_dict(tds)
    assert row["30 Day Price Change"] == "10.0; 11.0; 12.0"


def test_row_to_dict_sentiment_is_formatted():
    soup = BeautifulSoup(_synth_row_html(), "html.parser")
    tds = soup.find("tr").find_all("td")
    row = row_to_dict(tds)
    assert row["Bullish or Bearish"] == "Bull 50% / Neutral 30% / Bear 20%"


# =====================================================================
# End-to-end against the real v3 + v4 docx (hard acceptance: PK match)
# =====================================================================

def _read_csv(p: Path) -> list[dict]:
    with open(p, encoding="utf-8") as f:
        return list(csv.DictReader(f))


@pytest.mark.parametrize("docx_name,expected_rows", [
    ("biotech_catalysts_v3.docx", 600),
    ("biotech_catalysts_v4.docx", 100),
])
def test_real_docx_conversion_round_trip(docx_name, expected_rows, tmp_path):
    src = PROJECT_ROOT / "_csv_source" / docx_name
    if not src.exists():
        pytest.skip(f"{docx_name} not present in this environment")
    out = tmp_path / docx_name.replace(".docx", ".csv")
    stats = convert(src, out)
    assert stats.rows_written == expected_rows
    assert stats.rows_skipped == 0

    rows = _read_csv(out)
    assert len(rows) == expected_rows

    # Acceptance: no row has an empty Ticker (the primary identifier).
    assert all(r["Ticker"] for r in rows)
    # Every row's Catalyst Date is either blank or ISO format.
    for r in rows:
        cd = r["Catalyst Date"]
        if cd:
            assert len(cd) == 10 and cd[4] == "-" and cd[7] == "-", \
                f"Catalyst Date {cd!r} is not ISO YYYY-MM-DD"
    # Drug column never contains FDA-badge fragments like " FTD" / " BTD" / " ODD".
    for r in rows:
        for tag in (" FTD", " BTD", " ODD", " View Clinical"):
            assert tag not in r["Drug"], f"Drug {r['Drug']!r} contains badge fragment {tag!r}"


def test_auto_convert_directory_skips_when_csv_is_newer(tmp_path):
    """Auto-mode should skip files whose CSV is already newer than the docx."""
    # Create a fake .docx with HTML content
    import docx as docx_module
    fake_docx = tmp_path / "test.docx"
    d = docx_module.Document()
    d.add_paragraph("<table><tbody><tr>" + "<td></td>" * 20 + "</tr></tbody></table>")
    d.save(fake_docx)

    # Create a sibling .csv newer than the docx
    fake_csv = tmp_path / "test.csv"
    fake_csv.write_text("placeholder")
    import os
    os.utime(fake_csv, (fake_docx.stat().st_mtime + 100,
                        fake_docx.stat().st_mtime + 100))

    results = auto_convert_directory(tmp_path)
    assert results == []
    # CSV not touched
    assert fake_csv.read_text() == "placeholder"


def test_auto_convert_directory_force_overrides_freshness(tmp_path):
    import docx as docx_module
    fake_docx = tmp_path / "test.docx"
    d = docx_module.Document()
    # Minimal valid HTML with 20 cells
    cells_html = "".join(f"<td><div blurred-text='v{i}'>v{i}</div></td>" for i in range(20))
    d.add_paragraph(f"<table><tbody><tr>{cells_html}</tr></tbody></table>")
    d.save(fake_docx)
    fake_csv = tmp_path / "test.csv"
    fake_csv.write_text("stale,placeholder")
    import os
    os.utime(fake_csv, (fake_docx.stat().st_mtime + 100,
                        fake_docx.stat().st_mtime + 100))

    results = auto_convert_directory(tmp_path, force=True)
    assert len(results) == 1
    # CSV was rewritten
    assert "stale,placeholder" not in fake_csv.read_text()
