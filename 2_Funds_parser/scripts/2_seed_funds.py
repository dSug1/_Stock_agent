"""Seed the funds table from Input/list_of_funds.xlsx.

Reads the "Confirmed Funds" sheet. For each row where column A
(header "Selected funds") is non-empty, upserts into `funds`:
    name       <- column A (user-preferred display name)
    legal_name <- column B (SEC legal name)
    cik        <- column C (10-digit zero-padded)

Idempotent: re-running updates legal_name/name on CIK conflict.

Run from the repo root:
    python 2_Funds_parser/scripts/2_seed_funds.py
or from 2_Funds_parser/ with PYTHONPATH=src (matches the .bat
layout):
    python scripts\\2_seed_funds.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve()
PROJECT_ROOT = HERE.parents[1]
SRC = PROJECT_ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from database.db import get_connection, now_iso  # noqa: E402

XLSX_PATH = PROJECT_ROOT / "Input" / "list_of_funds.xlsx"
SHEET_NAME = "Confirmed Funds"


def read_selected_funds() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    rows = []
    for r in range(2, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        c = ws.cell(row=r, column=3).value
        if a is None or not str(a).strip():
            continue
        rows.append({
            "name": str(a).strip(),
            "legal_name": str(b).strip() if b else "",
            "cik": str(c).strip() if c else "",
        })
    return rows


def seed(conn, funds: list[dict]) -> tuple[int, int]:
    """Upsert each fund by CIK. Returns (inserted, updated)."""
    ts = now_iso()
    inserted = 0
    updated = 0
    for f in funds:
        existing = conn.execute(
            "SELECT id FROM funds WHERE cik = ?", (f["cik"],)
        ).fetchone()
        if existing is None:
            conn.execute(
                "INSERT INTO funds "
                "(cik, name, legal_name, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?)",
                (f["cik"], f["name"], f["legal_name"], ts, ts),
            )
            inserted += 1
        else:
            conn.execute(
                "UPDATE funds SET "
                " name = ?, legal_name = ?, updated_at = ? "
                "WHERE cik = ?",
                (f["name"], f["legal_name"], ts, f["cik"]),
            )
            updated += 1
    conn.commit()
    return inserted, updated


def main() -> int:
    funds = read_selected_funds()
    print(f"Read {len(funds)} selected funds from {XLSX_PATH.name}")

    if not funds:
        print("No rows selected in column A; nothing to do.")
        return 0

    missing = [
        f for f in funds if not f["cik"] or not f["legal_name"]
    ]
    if missing:
        print(f"[FATAL] {len(missing)} selected rows missing cik or legal_name:")
        for f in missing:
            print(f"  {f}")
        return 1

    conn = get_connection()
    try:
        inserted, updated = seed(conn, funds)
        total = conn.execute(
            "SELECT COUNT(*) AS n FROM funds"
        ).fetchone()["n"]
    finally:
        conn.close()

    print(
        f"Seed complete: {inserted} inserted, {updated} updated. "
        f"funds table now has {total} rows."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
